from django.shortcuts import render, get_object_or_404, redirect
from .models import StudentFee, TeacherSalary
from .forms import StudentFeeForm, TeacherSalaryForm
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.core.mail import send_mail, EmailMessage
from django.conf import settings
from .models import PaymentNotification
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_POST
from django.http import JsonResponse, FileResponse, Http404, HttpResponse
from django.utils import timezone
from django.utils.html import format_html
from .utils import generate_qr_code_base64, generate_receipt_pdf_weasyprint, generate_payslip_pdf_weasyprint
import razorpay
import json
from django.utils.html import strip_tags
from django.template.loader import render_to_string
from xhtml2pdf import pisa
import os
from weasyprint import HTML
from django.core.files.base import ContentFile
import openpyxl
from .forms import TeacherSalaryForm
from django.http import HttpResponseForbidden
# Create your views here.

client = razorpay.Client(auth=(settings.RAZORPAY_KEY_ID, settings.RAZORPAY_KEY_SECRET))

def is_accountant(user):
    return user.is_authenticated and user.role == 'accountant'

@login_required
def accountant_dashboard(request):
    if request.user.role != 'accountant':
        messages.error(request, "You are not authorized to perform this action.")
        return redirect('accountant_login')

    notifications = PaymentNotification.objects.order_by('-created_at')[:20]
    PaymentNotification.objects.filter(seen=False).update(seen=True)

    return render(request, 'accountant/dashboard.html', {
        'notifications': notifications,
    })

@login_required
def receive_student_fee(request, fee_id):
    fee = get_object_or_404(StudentFee, id=fee_id)

    if request.user.role != 'accountant':
        return redirect('accountant_login')

    fee.status = 'paid'
    fee.save()

    generate_receipt_pdf_weasyprint(fee)

    PaymentNotification.objects.create(
        student=fee.student,
        amount=fee.amount,
        status='success',
        message=f"Fee payment of ₹{fee.amount} received successfully on {fee.date}.",
        seen=False,
    )

    subject = "Fee Payment Received - Sree Vidyalaya EM High School"
    plain_message = "We have received your fee of ₹{} on {}. Thank you.".format(fee.amount, fee.date)

    html_message = format_html("""
        <div style="font-family: Arial, sans-serif; color: #333;">
            <h2 style="color: #2c3e50;">Payment Confirmation</h2>
            <p>Dear <strong>{name}</strong>,</p>
            <p>
                We have successfully received your payment of <strong>₹{amount}</strong> on 
                <strong>{date}</strong>.
            </p>
            <p>A receipt has been generated and can be accessed from your student portal.</p>
            <p style="margin-top: 20px;">Thank you for your timely payment.<br>
            <em>Sree Vidyalaya EM High School</em></p>
            <hr>
            <small>This is an automated email. Please do not reply to it.</small>
        </div>
    """, name=fee.student.get_full_name(), amount=fee.amount, date=fee.date)

    send_mail(
        subject=subject,
        message=plain_message,
        from_email=settings.DEFAULT_FROM_EMAIL,
        recipient_list=[fee.student.email],
        html_message=html_message,
    )

    messages.success(request, 'Fee marked as paid and PDF receipt generated.')
    return redirect('student_fee_list')

# ------------------- STUDENT FEES ------------------- #

@login_required
def student_fee_list(request):
    if request.user.role != 'accountant':
        return redirect('accountant_login')

    fees = StudentFee.objects.all().order_by('-date')
    return render(request, 'accountant/student_fee_list.html', {
        'fees': fees,
        'paid_fees': fees.filter(status='paid'),
        'unpaid_fees': fees.filter(status='unpaid'),
    })

@login_required
def student_fee_create(request):
    if request.user.role != 'accountant':
        return redirect('accountant_login')
    form = StudentFeeForm(request.POST or None)
    if form.is_valid():
        fee = form.save()
        subject = "🧾 New Fee Assigned - Sree Vidyalaya EM High School"
        message = (
            f"Dear {fee.student.get_full_name()},\n\n"
            f"A new fee of ₹{fee.amount} has been assigned to your account."
        )  # Fallback plain-text message

        html_message = format_html("""
            <div style="font-family: Arial, sans-serif; color: #333;">
                <h2 style="color: #2c3e50;">New Fee Assigned</h2>
                <p>Dear <strong>{name}</strong>,</p>
                <p>
                    A new fee of <strong>₹{amount}</strong> has been assigned to your account 
                    on <strong>{date}</strong>.
                </p>
                <p>Please log in to the student portal to view and pay your fee.</p>
                <p style="margin-top: 20px;">Thank you,<br>
                <em>Sree Vidyalaya EM High School</em></p>
                <hr>
                <small>This is an automated message. Please do not reply.</small>
            </div>
        """, name=fee.student.get_full_name(), amount=fee.amount, date=fee.date)

        send_mail(
            subject=subject,
            message=message,  # Plain-text fallback
            from_email=settings.DEFAULT_FROM_EMAIL,
            recipient_list=[fee.student.email],
            html_message=html_message  # HTML content
        )
        messages.success(request, 'Student fee created successfully')
        return redirect('student_fee_list')
    return render(request, 'accountant/create_student_fee.html', {'form': form})


@login_required
def student_fee_update(request, pk):
    fee = get_object_or_404(StudentFee, pk=pk)
    form = StudentFeeForm(request.POST or None, instance=fee)
    if form.is_valid():
        form.save()
        messages.success(request, 'Student fee updated')
        return redirect('student_fee_list')
    
    return render(request, 'accountant/student_fee_form.html', {'form': form})

@login_required
def student_fee_delete(request, pk):
    fee = get_object_or_404(StudentFee, pk=pk)
    fee.delete()
    messages.warning(request, 'Student fee deleted')
    return redirect('student_fee_list')

@login_required
def view_fee_receipt(request, fee_id):
    fee = get_object_or_404(StudentFee, id=fee_id)
    if not fee.receipt_pdf:
        messages.warning(request, "No receipt available for this payment.")
        return redirect('student_fee_list')

    return FileResponse(fee.receipt_pdf.open(), content_type='application/pdf')

@login_required
def export_fees_excel(request):
    if request.user.role != 'accountant':
        return redirect('accountant_login')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Student Fees"
    ws.append(['Student', 'Email', 'Amount', 'Status', 'Date', 'Receipt No'])

    for fee in StudentFee.objects.all():
        ws.append([
            fee.student.get_full_name(),
            fee.student.email,
            str(fee.amount),
            fee.status,
            fee.date.strftime('%Y-%m-%d'),
            fee.receipt_number or ''
        ])

    response = HttpResponse(content_type='application/vnd.ms-excel')
    response['Content-Disposition'] = 'attachment; filename=student_fees.xlsx'
    wb.save(response)
    return response
# ------------------- TEACHER SALARIES ------------------- #

@login_required
def teacher_salary_list(request):
    if request.user.role != 'accountant':
        return redirect('accountant_login')
    salaries = TeacherSalary.objects.all()
    return render(request, 'accountant/teacher_salary_list.html', {'salaries': salaries})

@login_required
def teacher_salary_create(request):
    form = TeacherSalaryForm(request.POST or None)
    if form.is_valid():
        form.save()
        messages.success(request, "Teacher salary entry created successfully.")
        return redirect('teacher_salary_list')
    return render(request, 'accountant/create_teacher_salary.html', {'form': form})

@login_required
def teacher_salary_update(request, pk):
    salary = get_object_or_404(TeacherSalary, pk=pk)
    form = TeacherSalaryForm(request.POST or None, instance=salary)
    if form.is_valid():
        form.save()
        messages.success(request, "Teacher salary updated successfully.")
        return redirect('teacher_salary_list')
    return render(request, 'accountant/teacher_salary_form.html', {'form': form})

@login_required
def teacher_salary_delete(request, pk):
    salary = get_object_or_404(TeacherSalary, pk=pk)
    salary.delete()
    messages.warning(request, "Teacher salary entry deleted.")
    return redirect('teacher_salary_list')

def generate_payslip_pdf(salary):
    template_path = 'accountant/payslip_template.html'
    context = {'salary': salary}
    html = render_to_string(template_path, context)
    
    result = io.BytesIO()
    pdf = pisa.pisaDocument(io.BytesIO(html.encode("UTF-8")), result)

    if not pdf.err:
        return ContentFile(result.getvalue())
    return None

@login_required
def pay_teacher_salary(request):
    if request.user.role != 'accountant':
        return redirect('accountant_login')

    if request.method == 'POST':
        form = TeacherSalaryForm(request.POST)
        if form.is_valid():
            salary = form.save(commit=False)
            salary.payment_method = 'cash'
            salary.paid_date = timezone.now().date()
            salary.save()
            send_payslip_email(salary)

            # Generate PDF
            generate_payslip_pdf_weasyprint(salary)

            # Email to teacher
            if salary.teacher.email:
                subject = f"Payslip for {salary.month} {salary.year}"
                message = f"Dear {salary.teacher.get_full_name()},\n\nYour salary of ₹{salary.amount} has been paid.\nPayslip is attached.\n\nRegards,\nAdmin"
                email = EmailMessage(subject, message, settings.DEFAULT_FROM_EMAIL, [salary.teacher.email])
                if salary.payslip:
                    email.attach_file(salary.payslip.path)
                email.send()

            messages.success(request, "Salary paid in cash, payslip generated, and email sent.")
            return redirect('teacher_salary_list')
    else:
        form = TeacherSalaryForm()

    return render(request, 'accountant/teacher_salary_form.html', {'form': form})


def send_email_to_teacher(salary):
    subject = f"Salary Credited: {salary.month}"
    html = render_to_string('salary_credited.html', {'salary': salary})
    send_mail(subject, strip_tags(html), settings.DEFAULT_FROM_EMAIL,
              [salary.teacher.email], html_message=html)

@login_required
def initiate_teacher_salary_payment(request, salary_id):
    salary = get_object_or_404(TeacherSalary, id=salary_id)

    if request.user.role != 'accountant':
        return redirect('accountant_login')

    # Initialize Razorpay client
    client = razorpay.Client(auth=(settings.RAZORPAY_KEY_ID, settings.RAZORPAY_KEY_SECRET))

    # Create Razorpay Order
    payment = client.order.create({
        'amount': int(salary.amount * 100),  # Amount in paise
        'currency': 'INR',
        'payment_capture': '1',
        'notes': {
            'teacher_name': salary.teacher.get_full_name(),
            'month': salary.month,
            'year': salary.year
        }
    })

    # Save order ID
    salary.payment_id = payment['id']
    salary.payment_method = 'online'
    salary.save()
    messages.info(request, f"Initiated salary payment for {salary.teacher.get_full_name()}.")
    return render(request, 'accountant/teacher_salary_pay.html', {
        'salary': salary,
        'razorpay_key': settings.RAZORPAY_KEY_ID,
        'order_id': payment['id'],
    })

@csrf_exempt
@require_POST
@login_required
def teacher_salary_payment_success(request):
    try:
        data = json.loads(request.body)
        razorpay_order_id = data.get('razorpay_order_id')
        razorpay_payment_id = data.get('razorpay_payment_id')
        razorpay_signature = data.get('razorpay_signature')
        salary_id = data.get('salary_id')

        salary = get_object_or_404(TeacherSalary, id=salary_id)

        # Verify Razorpay signature
        client.utility.verify_payment_signature({
            'razorpay_order_id': razorpay_order_id,
            'razorpay_payment_id': razorpay_payment_id,
            'razorpay_signature': razorpay_signature,
        })

        salary.payment_id = razorpay_payment_id
        salary.paid_date = timezone.now().date()
        salary.payment_method = 'online'
        salary.save()

        salary.generate_payslip_pdf()
        salary.save()

        # TODO: Send email to teacher (will add in next step)

        return JsonResponse({'success': True})

    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)
    
@login_required
def download_teacher_payslip(request, pk):
    salary = get_object_or_404(TeacherSalary, pk=pk)

    # Authorization check
    if request.user != salary.teacher and request.user.role != 'accountant':
        return HttpResponseForbidden("Unauthorized")

    # File check
    if not salary.payslip or not salary.payslip.path or not os.path.exists(salary.payslip.path):
        messages.warning(request, "Payslip file not found.")
        if request.user.role == 'accountant':
            return redirect('teacher_salary_list')
        return redirect('my_payslips')

    # Return file response with download header
    response = FileResponse(
        salary.payslip.open('rb'),
        content_type='application/pdf'
    )
    response['Content-Disposition'] = f'attachment; filename=payslip_{salary.teacher.id}_{salary.month}_{salary.year}.pdf'
    return response

def send_payslip_email(salary):
    if not salary.payslip:
        return

    subject = f"Payslip for {salary.month}, {salary.year}"
    body = f"Dear {salary.teacher.get_full_name()},\n\nPlease find attached your payslip for the month of {salary.month} {salary.year}.\n\nRegards,\nAdmin"
    email = EmailMessage(
        subject=subject,
        body=body,
        from_email=settings.DEFAULT_FROM_EMAIL,
        to=[salary.teacher.email],
    )
    email.attach_file(salary.payslip.path)
    email.send()
    
@login_required
def export_teacher_salaries_excel(request):
    if request.user.role != 'accountant':
        return HttpResponse("Unauthorized", status=403)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Teacher Salaries"
    ws.append(["Teacher", "Email", "Amount", "Month", "Year", "Paid Date", "Method", "Note"])

    for salary in TeacherSalary.objects.all():
        ws.append([
            salary.teacher.get_full_name(),
            salary.teacher.email,
            str(salary.amount),
            salary.month,
            salary.year,
            salary.paid_date.strftime("%Y-%m-%d"),
            salary.payment_method,
            salary.note or "",
        ])

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="teacher_salaries.xlsx"'
    wb.save(response)
    return response